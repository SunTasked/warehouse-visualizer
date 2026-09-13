#!/usr/bin/env python3
"""Import the CML plant's buildings from its workbook (sheet PLAN_Sortie).

Run:  python scripts/import_plan_xlsx.py
      python scripts/import_plan_xlsx.py --building "BATIMENT 13A" "BATIMENT 12B"
Then: python scripts/check_plan.py schema/CML.plan.json

--------------------------------------------------------------------------
How the spreadsheet encodes a warehouse
--------------------------------------------------------------------------
The sheet is a floor plan drawn with cells. Reading it does *not* require
guessing from geometry, because the drawing is self-describing:

* Every storage location has a **label cell** (e.g. 'AA01') written *inside
  the aisle that serves it*, on the line nearest the rack it names. So for
  aisle A, the 'AA01, AA03, ...' line names the rack above it and the
  'AA02, AA04, ...' line the rack below. Odd codes therefore always face
  south and even codes north (verified: true for every two-sided aisle in
  13A, no exceptions). 16G is drawn the other way round: its aisles run
  north-south, with the labels in columns and every rack facing east or west.

* Every drawn box carries a formula that identifies it:
    =VLOOKUP(<its own label cell>, Analyse!...)  -> the FRONT position, the
                                                    one touching the aisle
    =<another box cell>                          -> the same location, one
                                                    position further back
  Following those references gives each location its exact set of cells, and
  the count of them is the rack's depth in pallet positions. A location is
  one addressable lane: depth N, LEVELS high.

* 38 boxes in 13A carry a broken '=VLOOKUP(,...)' (the reference was deleted
  at some point). Their label is recovered from the adjacent label cell.

* Black-filled boxes are drawn but carry no formula and no code - blocked
  positions (pillars and the like). They become holes in the racking.

* A box whose label cell is blank (07D's would-be DM02, 06F's would-be FP12,
  15K's T358) names nothing: it is left out as a blocked position too, per
  the plant owner.

* Aisles P and D of 06F are drawn one row tall with a single line of even
  codes, and the lanes on *both* sides look up the same label: FP04 is one
  address covering two facing one-deep lanes (the picking history only ever
  uses the plain code). Such a code becomes two slots, suffixed A for the
  lane north of (or west of) the aisle and B for the other, per the plant
  owner.

* Areas forklifts can't enter are named on the plan - some as merged cells
  ("Lithium"), most as drawn text boxes, which openpyxl doesn't load and are
  read from the drawing XML ("TRAIN", "ZONE MU"). Per the plant owner they
  become inaccessible zones. A text box is the area itself, except "ZONE MU",
  which only labels the medium-bordered box drawn around it.

* Bold borders mark back-to-back facings, as the plan's author intended, but
  the file is not consistent about it (the D/E block uses a thin line for the
  same thing). They are therefore *not* used here - the formulas settle every
  case on their own.

--------------------------------------------------------------------------
How it becomes a warehouse
--------------------------------------------------------------------------
One Excel location -> one Slot; its depth cells -> that slot's sub-slots;
LEVELS -> pallet tiers per sub-slot (content, not config, so it is only a
capacity here).

The drawing is schematic, not to scale: a one-deep lane is often drawn two
cells tall while a three-deep lane gets three cells. So cell *extents* are
not used as distances. Instead every column gets a fixed pitch and every row
a pitch depending on its role (rack row vs aisle line), and a slot is
anchored at the aisle-facing edge of its front box. Relative positions - the
thing the plan actually encodes - are preserved exactly.

Aisles come out of the same label lines: the corridor centreline for aisle A
is the line its 'AA..' labels are written on, spanning the labels. One aisle
letter can label two separate aisles (10H's short east end of aisle A);
each becomes its own corridor. Aisles are joined by a trunk in the empty band
between the rack blocks - a column band where aisles run east-west, 16G's
cross-aisle row band where they run north-south - and a short spur links an
aisle that can't reach the trunk to a neighbouring one that does.

--------------------------------------------------------------------------
How the buildings become a plant
--------------------------------------------------------------------------
The workbook draws every building in the same template and says nothing
about where they stand. Per the plant owner, they are stacked north to south
in sheet order (the first --building furthest north), west walls aligned,
ROAD_GAP metres apart, and linked along the central aisle: each trunk runs on
to a door in the building's south wall, and a short connector crosses to a
door in the north wall of the building below (16G, with no central aisle
north-south, is entered down its aisle nearest the one above).

Every building gets a lift station and a delivery space on the clear floor
south of its last rack, beside the corridor that runs along that strip; the
first building's lift station is where the forklift starts. Neither is in
the workbook.
"""
import argparse
import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# --- Metric assumptions (the drawing carries no scale) ----------------------
LANE_W = 1.2        # one location's width along the rack face
POS_D = 1.2         # one pallet position's depth into the rack
AISLE_W = 3.5       # a normal picking aisle
WIDE_AISLE_W = 5.0  # a main cross-aisle (drawn 4+ rows tall, e.g. aisle C)
WIDE_AISLE_ROWS = 4
MIN_FREE_ROW = 0.3  # a rack-free aisle row's pitch, however little floor is left to share
LEVELS = 3          # pallet tiers per position
MARGIN = 1.0        # clear floor between the outermost racking and the wall

# --- Circulation and facilities ---------------------------------------------
PATH_W = 2.2        # corridor width drawn in the app
CLEAR = 0.5         # clearance between a corridor or pad and racking or a wall
PAD_W, PAD_D = 6.0, 2.0  # lift station / delivery space footprint (the app's fixed size)
PAD_GAP = 0.4       # between the dock corridor and the pads beside it
ROAD_GAP = 10.0     # clear ground between two stacked buildings

DEFAULT_BUILDINGS = [
    "BATIMENT 13A", "BATIMENT 12B", "BATIMENT 08C", "BATIMENT 07D", "BATIMENT 06F",
    "BATIMENT 10H", "BATIMENT 14J", "BATIMENT 15K", "BATIMENT 16G",
]

# --- Areas forklifts can't enter ---------------------------------------------
# Plan wording (matched as whole words, any case) -> the zone's label.
ZONE_LABELS = {"TRAIN": "Train", "ZONE MU": "Zone MU", "LITHIUM": "Lithium"}
# Labels that name a bordered area rather than being the area.
BOXED_ZONES = {"Zone MU"}
COVER = 0.4  # a shape's edge cell counts once the shape covers this much of it

# rotationDeg -> facing, per Slots.tsx's rotation math (see the example
# generator, which documents the same mapping).
FACING = {"N": 0, "W": 90, "S": 180, "E": 270}

VLOOKUP = re.compile(r"^=VLOOKUP\(\$?([A-Z]{1,2})\$?(\d+)\s*,", re.I)
BROKEN = re.compile(r"^=VLOOKUP\(\s*,", re.I)
REF = re.compile(r"^=\$?([A-Z]{1,2})\$?(\d+)$", re.I)
CODE = re.compile(r"^[A-Z]{2}\d{2,3}$")
ANCHOR = re.compile(
    r"<xdr:(from|to)><xdr:col>(\d+)</xdr:col><xdr:colOff>(-?\d+)</xdr:colOff>"
    r"<xdr:row>(\d+)</xdr:row><xdr:rowOff>(-?\d+)</xdr:rowOff></xdr:\1>"
)


# ---------------------------------------------------------------------------
# Reading the sheet
# ---------------------------------------------------------------------------
def _style(side):
    return side.style if side is not None and side.style else None


def _is_black(cell):
    f = cell.fill
    return (
        f is not None
        and f.patternType == "solid"
        and f.fgColor is not None
        and f.fgColor.type == "theme"
        and f.fgColor.theme == 1
    )


def _is_boxed(cell):
    b = cell.border
    return (
        any(_style(getattr(b, s)) for s in ("left", "right", "top", "bottom"))
        or _is_black(cell)
        or (isinstance(cell.value, str) and cell.value.startswith("="))
    )


def _label_at(ws, r, c):
    v = ws.cell(r, c).value
    return v.strip() if isinstance(v, str) and CODE.match(str(v).strip()) else None


def _resolve(ws, r, c, seen=None):
    """-> (location code, label cell, front box) for the box at (r, c).

    The front box is the VLOOKUP box a back position's reference chain ends
    at - what tells apart the two lanes sharing one code. None when a back
    position points straight at the label cell instead.
    """
    seen = seen or set()
    if (r, c) in seen:
        return None, None, None
    seen.add((r, c))
    v = ws.cell(r, c).value
    if not isinstance(v, str) or not v.startswith("="):
        return None, None, None
    m = VLOOKUP.match(v)
    if m:
        lr, lc = int(m.group(2)), column_index_from_string(m.group(1))
        return _label_at(ws, lr, lc), (lr, lc), (r, c)
    if BROKEN.match(v):
        for dr, dc in ((-1, 0), (1, 0), (0, -1), (0, 1)):
            lab = _label_at(ws, r + dr, c + dc)
            if lab:
                return lab, (r + dr, c + dc), (r, c)
        return None, None, None
    m = REF.match(v)
    if m:
        tr, tc = int(m.group(2)), column_index_from_string(m.group(1))
        lab = _label_at(ws, tr, tc)
        if lab:  # points straight at the label cell
            return lab, (tr, tc), None
        return _resolve(ws, tr, tc, seen)
    return None, None, None


def find_building_rows(ws, name):
    """Interior row range of a building, from the rotated markers in column A.

    The north and south walls are drawn as medium-weight borders spanning the
    block; the floor is everything between them - including the clear strips
    with no racking, which is where the dock ends up.
    """
    markers = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().upper().startswith("BATIMENT"):
            markers.append((r, v.strip()))
    for i, (r, label) in enumerate(markers):
        if label.upper() != name.upper():
            continue
        end = markers[i + 1][0] - 1 if i + 1 < len(markers) else ws.max_row
        walls = [
            rr
            for rr in range(r, end + 1)
            if sum(
                1
                for c in range(1, 120)
                if _style(ws.cell(rr, c).border.top) == "medium" or _style(ws.cell(rr, c).border.bottom) == "medium"
            )
            > 20
        ]
        if len(walls) >= 2:
            return walls[0] + 1, walls[-1] - 1, label
        return r + 1, end - 1, label
    raise SystemExit(f"building {name!r} not found; have: {[m[1] for m in markers]}")


def _sign(v):
    return (v > 0) - (v < 0)


def _lane(cells, front, label):
    """One addressable lane: its cells front-to-back, facing, label cell."""
    (fr, fc), (lr, lc) = front, label
    dr, dc = lr - fr, lc - fc
    if dr and dc:
        return None
    face = "N" if dr < 0 else "S" if dr > 0 else "W" if dc < 0 else "E"
    # Depth runs away from the aisle, i.e. opposite the facing.
    step = {"N": (1, 0), "S": (-1, 0), "W": (0, 1), "E": (0, -1)}[face]
    ordered = sorted(cells, key=lambda rc: (rc[0] - fr) * step[0] + (rc[1] - fc) * step[1])
    return {"cells": ordered, "face": face, "label": label}


def extract(ws, r0, r1, c1=200):
    """-> (locations, blacks, shared). Each location: cells front-to-back, facing, label cell."""
    boxes, blacks, blank = {}, [], []
    for r in range(r0, r1 + 1):
        for c in range(1, c1 + 1):
            cell = ws.cell(r, c)
            if not _is_boxed(cell):
                continue
            if _is_black(cell):
                blacks.append((r, c))
                continue
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue  # blank continuation cell of a box drawn several cells tall
            code, lab, front = _resolve(ws, r, c)
            if not code:
                blank.append((r, c))
                continue
            boxes[(r, c)] = (code, lab, front)
    if blank:
        print(f"  - {len(blank)} boxes with a blank label, left out as blocked: "
              f"{[f'{get_column_letter(c)}{r}' for r, c in blank]}")
        blacks.extend(blank)

    grouped = defaultdict(list)
    for cell, (code, lab, front) in boxes.items():
        grouped[code].append((cell, lab, front))

    locations, shared = {}, []
    for code, members in sorted(grouped.items()):
        fronts = sorted({front for cell, _, front in members if front == cell})
        labels = {lab for _, lab, _ in members}
        if len(labels) != 1:
            print(f"  ! {code}: written in {len(labels)} places, skipped")
            continue
        label = labels.pop()

        if len(fronts) == 1:
            lane = _lane([cell for cell, _, _ in members], fronts[0], label)
            if lane is None:
                print(f"  ! {code}: label is diagonal from its front box, skipped")
                continue
            locations[code] = lane
            continue

        # One code, two lanes facing each other across its label line.
        sides = [(_sign(f[0] - label[0]), _sign(f[1] - label[1])) for f in fronts]
        if len(fronts) != 2 or sides[0] != (-sides[1][0], -sides[1][1]):
            print(f"  ! {code}: {len(fronts)} front cells, skipped")
            continue
        for suffix, front, side in zip("AB", fronts, sides):  # sorted: A is the north (or west) lane
            cells = [
                cell
                for cell, _, f in members
                if f == front
                or (f is None and (_sign(cell[0] - label[0]), _sign(cell[1] - label[1])) == side)
            ]
            lane = _lane(cells, front, label)
            if lane is None:
                print(f"  ! {code}{suffix}: label is diagonal from its front box, skipped")
                continue
            locations[code + suffix] = lane
        shared.append(code)
    return locations, blacks, shared


# ---------------------------------------------------------------------------
# Inaccessible zones
# ---------------------------------------------------------------------------
def _attrs(tag):
    return dict(re.findall(r'([\w:]+)="([^"]*)"', tag))


def _rels(z, path):
    try:
        xml = z.read(path).decode("utf8")
    except KeyError:
        return {}
    return {a["Id"]: a["Target"] for a in map(_attrs, re.findall(r"<Relationship\b[^>]*>", xml))}


def _drawing_xml(workbook_path, sheet_name):
    """The drawing layer (shapes, text boxes) of one sheet, which openpyxl doesn't load."""
    with zipfile.ZipFile(workbook_path) as z:
        sheets = map(_attrs, re.findall(r"<sheet\b[^>]*>", z.read("xl/workbook.xml").decode("utf8")))
        rid = next((s.get("r:id") for s in sheets if s.get("name") == sheet_name), None)
        target = _rels(z, "xl/_rels/workbook.xml.rels").get(rid)
        if not target:
            return ""
        sheet_path = "xl/" + target.lstrip("/").removeprefix("xl/")
        folder, base = sheet_path.rsplit("/", 1)
        drawings = [t.rsplit("/", 1)[1] for t in _rels(z, f"{folder}/_rels/{base}.rels").values() if "drawings/" in t]
        return "".join(z.read(f"xl/drawings/{d}").decode("utf8", "replace") for d in drawings if d.endswith(".xml"))


def _zone_label(text):
    if not isinstance(text, str):
        return None
    words = " ".join(text.upper().split())
    return next((name for key, name in ZONE_LABELS.items() if re.search(rf"\b{re.escape(key)}\b", words)), None)


def _col_emu(ws, c):
    dim = ws.column_dimensions.get(get_column_letter(c))
    width = dim.width if dim is not None and dim.width else (ws.sheet_format.defaultColWidth or 8.43)
    return (width * 7 + 5) * 9525


def _row_emu(ws, r):
    dim = ws.row_dimensions.get(r)
    height = dim.height if dim is not None and dim.height else (ws.sheet_format.defaultRowHeight or 15)
    return height * 12700


def _enclosing_box(ws, cells, reach=60):
    """The medium-bordered box drawn around a label, as cells, or None."""
    r0, c0, r1, c1 = cells

    def medium(cell, side):
        return _style(getattr(cell.border, side)) == "medium"

    top = next((r for r in range(r0, max(1, r0 - reach), -1)
                if medium(ws.cell(r, c0), "top") or medium(ws.cell(r - 1, c0), "bottom")), None)
    bottom = next((r for r in range(r1, r1 + reach)
                   if medium(ws.cell(r, c0), "bottom") or medium(ws.cell(r + 1, c0), "top")), None)
    left = next((c for c in range(c0, max(1, c0 - reach), -1)
                 if medium(ws.cell(r0, c), "left") or medium(ws.cell(r0, c - 1), "right")), None)
    right = next((c for c in range(c1, c1 + reach)
                  if medium(ws.cell(r0, c), "right") or medium(ws.cell(r0, c + 1), "left")), None)
    if None in (top, bottom, left, right):
        return None
    return top, left, bottom, right


def read_zones(ws, workbook_path, sheet_name):
    """-> [(label, (r0, c0, r1, c1))]: the sheet's no-access areas, in cells."""
    zones, seen = [], set()
    xml = _drawing_xml(workbook_path, sheet_name)
    for body in re.findall(r"<xdr:twoCellAnchor\b[^>]*>(.*?)</xdr:twoCellAnchor>", xml, re.S):
        label = _zone_label(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", body)))
        anchors = {m[0]: tuple(int(v) for v in m[1:]) for m in ANCHOR.findall(body)}
        if not label or len(anchors) != 2:
            continue
        c0, co0, r0, ro0 = anchors["from"]
        c1, co1, r1, ro1 = anchors["to"]
        c0, r0, c1, r1 = c0 + 1, r0 + 1, c1 + 1, r1 + 1  # the XML counts from 0
        cells = (
            r0 if 1 - ro0 / _row_emu(ws, r0) >= COVER else r0 + 1,
            c0 if 1 - co0 / _col_emu(ws, c0) >= COVER else c0 + 1,
            r1 if ro1 / _row_emu(ws, r1) >= COVER else r1 - 1,
            c1 if co1 / _col_emu(ws, c1) >= COVER else c1 - 1,
        )
        if (label, cells) in seen:  # the drawing repeats some of its shapes
            continue
        seen.add((label, cells))
        if label in BOXED_ZONES:
            cells = _enclosing_box(ws, cells) or cells
        zones.append((label, cells))
    for mr in ws.merged_cells.ranges:
        label = _zone_label(ws.cell(mr.min_row, mr.min_col).value)
        if label:
            zones.append((label, (mr.min_row, mr.min_col, mr.max_row, mr.max_col)))
    return zones


# ---------------------------------------------------------------------------
# Geometry
# ---------------------------------------------------------------------------
def _clusters(values, gap):
    """Sorted distinct values, split wherever two neighbours are more than `gap` apart."""
    out = []
    for v in sorted(set(values)):
        if out and v - out[-1][-1] <= gap:
            out[-1].append(v)
        else:
            out.append([v])
    return out


def ew_aisles(locations):
    """Aisles of east/west-facing racks: one per code prefix and cluster of label columns."""
    by_prefix = defaultdict(list)
    for code, loc in locations.items():
        if loc["face"] in "WE":
            by_prefix[code[:2]].append(loc)
    out = []
    for prefix, locs in sorted(by_prefix.items()):
        groups = _clusters([l["label"][1] for l in locs], 2)
        for n, cols in enumerate(groups, 1):
            members = [l for l in locs if cols[0] <= l["label"][1] <= cols[-1]]
            out.append({
                "id": f"aisle-{prefix}" + (f"-{n}" if len(groups) > 1 else ""),
                "cols": cols,
                "rows": sorted({l["label"][0] for l in members}),
            })
    return out


def ns_aisles(locations, blockers):
    """Aisles of north/south-facing racks: one per code prefix and run of label rows.

    Two label rows of one prefix are one aisle unless racking or a zone stands
    between them, across the columns their labels span - which is what makes
    10H's short east end of aisle A, behind a rack, an aisle of its own.
    """
    by_prefix = defaultdict(lambda: defaultdict(list))  # prefix -> label row -> label columns
    for code, loc in locations.items():
        if loc["face"] in "NS":
            by_prefix[code[:2]][loc["label"][0]].append(loc["label"][1])
    out = []
    for prefix, rows in sorted(by_prefix.items()):
        groups = []
        for r in sorted(rows):
            if groups:
                previous = groups[-1]
                cols = rows[r] + [c for pr in previous for c in rows[pr]]
                between = range(previous[-1] + 1, r)
                if not any((rr, c) in blockers for rr in between for c in range(min(cols), max(cols) + 1)):
                    previous.append(r)
                    continue
            groups.append([r])
        for n, group in enumerate(groups, 1):
            out.append({
                "id": f"aisle-{prefix}" + (f"-{n}" if len(groups) > 1 else ""),
                "rows": group,
                "cols": sorted({c for r in group for c in rows[r]}),
            })
    return out


def build_axes(locations, r0, r1, transposed=False, zone_cells=frozenset()):
    """Per-row heights and per-column widths.

    A row that carries the labels of a north/south-facing rack *is* that
    rack's aisle, so it gets aisle pitch; every other row is a rack row at one
    pallet position deep. Same idea transposed for east/west racks, whose
    labels sit in columns instead.
    """
    aisle_rows = {loc["label"][0] for loc in locations.values() if loc["face"] in "NS"}

    # An aisle drawn with blank rows between its two label lines (the main
    # cross-aisle) is wider; absorb those blanks into the group. A zone's rows
    # aren't blank: absorbed, 10H's Lithium store and Zone MU were squashed to
    # an aisle's share of height, and Lithium spilled over the racks beside it.
    if aisle_rows:
        lo, hi = min(aisle_rows), max(aisle_rows)
        occupied = {r for loc in locations.values() for r, _ in loc["cells"]} | {r for r, _ in zone_cells}
        for r in range(lo, hi + 1):
            if r not in aisle_rows and r not in occupied:
                if any(x in aisle_rows for x in range(lo, r)) and any(x in aisle_rows for x in range(r + 1, hi + 1)):
                    aisle_rows.add(r)

    groups, current = [], []
    for r in range(r0, r1 + 2):
        if r in aisle_rows:
            current.append(r)
        elif current:
            groups.append(current)
            current = []

    # An aisle line that also carries racking somewhere else along its length
    # (the plan has a few) must still be at least one position deep, or that
    # rack would overflow its band and collide with the one behind it.
    occupied_rows = {r for loc in locations.values() for r, _ in loc["cells"]}
    row_h = {r: POS_D for r in range(r0, r1 + 1)}
    for g in groups:
        total = WIDE_AISLE_W if len(g) >= WIDE_AISLE_ROWS else AISLE_W
        free = [r for r in g if r not in occupied_rows]
        forced = [r for r in g if r in occupied_rows]
        spare = total - len(forced) * POS_D
        for r in forced:
            row_h[r] = POS_D
        # Rows with no racking anywhere along them just share what's left,
        # however many of them the drawing spent: 07D's cross-aisle H is drawn
        # nine rows tall, and a position's pitch each made it 10.8 m wide -
        # wide enough that its own racks sat nearer the dock corridor behind
        # them than the aisle in front.
        for r in free:
            row_h[r] = max(MIN_FREE_ROW, spare / len(free))

    cols = [c for loc in locations.values() for _, c in loc["cells"]]
    c0, c1 = min(cols), max(cols)
    occupied_cols = set(cols)
    col_w = {c: LANE_W for c in range(c0 - 1, c1 + 2)}
    for aisle in ew_aisles(locations):
        if transposed:
            # 16G: an aisle's label columns, and anything drawn between them,
            # share one aisle's width; a column that carries racking elsewhere
            # (its two rack blocks don't line up) keeps a position's pitch.
            span = range(aisle["cols"][0], aisle["cols"][-1] + 1)
            for c in span:
                if c not in occupied_cols:
                    col_w[c] = max(col_w.get(c, LANE_W), AISLE_W / len(span))
        else:
            # A side block of east/west racks (13A's Z block): each label
            # column is an aisle's width, as it has always been imported.
            for c in aisle["cols"]:
                if c not in occupied_cols:
                    col_w[c] = AISLE_W

    return row_h, col_w, (c0 - 1, c1 + 1)


def axis_edges(sizes, lo, hi):
    """-> {index: (start, end)} accumulating from lo."""
    out, acc = {}, 0.0
    for i in range(lo, hi + 1):
        w = sizes.get(i, POS_D)
        out[i] = (acc, acc + w)
        acc += w
    return out, acc


def _widest_interior_band(indices, occupied_indices):
    """The widest run of consecutive empty indices with something occupied on both sides."""
    bands = []
    for i in indices:
        if bands and i == bands[-1][-1] + 1:
            bands[-1].append(i)
        else:
            bands.append([i])
    interior = [b for b in bands if any(i < b[0] for i in occupied_indices) and any(i > b[-1] for i in occupied_indices)]
    return max(interior, key=len) if interior else (bands[-1] if bands else None)


def import_building(ws, sheet_name, sheet_zones, north_link_x=None, south_link=False):
    """One building in local coordinates: its south-west wall corner at (0, 0).

    north_link_x: where the building to the north enters from (its south
    door's x), or None when there is none; south_link: whether a building
    follows to the south.
    """
    r0, r1, label = find_building_rows(ws, sheet_name)
    # Named after the sheet's own block without the "BATIMENT" prefix (so
    # "BATIMENT 13A" -> "13A"): the wall loop's id, and what every facility
    # and corridor inside it belongs to.
    building_id = re.sub(r"^BATIMENT\s+", "", label, flags=re.I).strip() or label
    print(f"{label}: rows {r0}..{r1}")
    locations, blacks, shared = extract(ws, r0, r1)
    racking = {(r, c) for loc in locations.values() for r, c in loc["cells"]}

    zones = []
    for zone_label, (zr0, zc0, zr1, zc1) in sheet_zones:
        rows, cols = (max(zr0, r0), min(zr1, r1)), (max(zc0, 2), zc1)
        if rows[0] > rows[1]:
            continue
        cells = {(r, c) for r in range(rows[0], rows[1] + 1) for c in range(cols[0], cols[1] + 1)}
        if cells & racking:
            print(f"  ! {zone_label} zone {get_column_letter(cols[0])}{rows[0]} covers racking, skipped")
            continue
        zones.append({"label": zone_label, "rows": rows, "cols": cols, "cells": cells})
    zone_cells = set().union(*(z["cells"] for z in zones)) if zones else set()
    blockers = racking | zone_cells      # what a corridor may not cross
    occupied = blockers | set(blacks)    # what a trunk band keeps entirely clear of

    transposed = sum(l["face"] in "WE" for l in locations.values()) > len(locations) / 2
    print(f"  {len(locations)} locations, {sum(len(l['cells']) for l in locations.values())} positions, "
          f"{len(blacks)} blocked" + (f", {len(shared)} codes split into A/B lanes" if shared else "")
          + (f", zones: {', '.join(z['label'] for z in zones)}" if zones else "")
          + (", aisles run north-south" if transposed else ""))

    row_h, col_w, (_, cmax) = build_axes(locations, r0, r1, transposed, zone_cells)
    if zones:
        cmax = max(cmax, max(z["cols"][1] for z in zones) + 1)
    # The floor is the building's whole interior, not just its racked part -
    # the clear strip south of the last rack is where the dock goes.
    rmin, rmax = r0, r1
    cmin = 2  # column 1 is the rotated building-name strip, outside the wall

    if transposed:
        # The cross-aisle between rack blocks is the trunk: one wide aisle,
        # however many rows it was drawn.
        occupied_rows = {r for r, _ in occupied}
        band_rows = _widest_interior_band(
            [r for r in range(rmin, rmax + 1) if r not in occupied_rows], occupied_rows)
        for r in band_rows:
            row_h[r] = max(MIN_FREE_ROW, WIDE_AISLE_W / len(band_rows))

    ys, total_h = axis_edges(row_h, rmin, rmax)
    xs, total_w = axis_edges(col_w, cmin, cmax)
    last_rack = max(r for r, _ in occupied)

    # The dock strip holds a corridor along the last rack and, south of it,
    # the pads. Where the drawn strip is too shallow for both, the building
    # grows south rather than a pad poking through the wall.
    rack_south = total_h - ys[last_rack][1] + MARGIN
    pad_bottom = rack_south - CLEAR - PATH_W - PAD_GAP - PAD_D
    extra_south = max(0.0, CLEAR - pad_bottom)
    base = MARGIN + extra_south

    # Excel rows run downward; world +Y runs up. Flip, and inset by MARGIN so
    # the outermost racking isn't flush against the wall.
    def y_top(r):
        return total_h - ys[r][0] + base

    def y_bot(r):
        return total_h - ys[r][1] + base

    def y_mid(r):
        return (y_top(r) + y_bot(r)) / 2

    def x_left(c):
        return xs[c][0] + MARGIN

    def x_right(c):
        return xs[c][1] + MARGIN

    def x_mid(c):
        return (x_left(c) + x_right(c)) / 2

    width = total_w + 2 * MARGIN
    height = total_h + 2 * MARGIN + extra_south
    print(f"  envelope {width:.1f} x {height:.1f} m" + (f" (grown {extra_south:.1f} m south for the dock)" if extra_south else ""))

    # --- slots -------------------------------------------------------------
    slots = []
    for code, loc in sorted(locations.items()):
        fr, fc = loc["cells"][0]
        face = loc["face"]
        depth = len(loc["cells"])
        if face == "S":      # opening faces south: front box's south edge is the aisle
            x, y = x_mid(fc), y_bot(fr) + POS_D / 2
        elif face == "N":
            x, y = x_mid(fc), y_top(fr) - POS_D / 2
        elif face == "W":
            x, y = x_left(fc) + POS_D / 2, y_mid(fr)
        else:
            x, y = x_right(fc) - POS_D / 2, y_mid(fr)
        slot = {"id": code, "x": x, "y": y}
        if FACING[face]:
            slot["rotationDeg"] = FACING[face]
        if depth > 1:
            slot["depth"] = depth
        slots.append(slot)

    # --- zones ---------------------------------------------------------------
    zone_counts = defaultdict(int)
    for z in zones:
        zone_counts[z["label"]] += 1
    zone_seen = defaultdict(int)
    zones_out = []
    for z in sorted(zones, key=lambda z: (z["rows"][0], z["cols"][0])):
        zone_seen[z["label"]] += 1
        slug = z["label"].lower().replace(" ", "-")
        x0, x1 = x_left(z["cols"][0]), x_right(z["cols"][1])
        yb, yt = y_bot(z["rows"][1]), y_top(z["rows"][0])
        zones_out.append({
            "id": f"{building_id}-{slug}" + (f"-{zone_seen[z['label']]}" if zone_counts[z["label"]] > 1 else ""),
            "label": z["label"],
            "x": (x0 + x1) / 2, "y": (yb + yt) / 2, "width": x1 - x0, "depth": yt - yb,
        })

    dock_y = y_bot(last_rack) - CLEAR - PATH_W / 2
    paths = []
    north_door = south_door = None

    if not transposed:
        # --- aisles run east-west ---------------------------------------------
        def clear_between_cols(rows, ca, cb):
            """No racking or zone between two columns, along a corridor's own rows.

            Blocked cells don't count: they are pillars, which a corridor already
            runs past within its own span (06F's aisle D has three), and one
            standing at the mouth of that aisle had cut it off from the trunk.
            """
            lo, hi = sorted((ca, cb))
            return not any((r, c) in blockers for r in rows for c in range(lo + 1, hi))

        # The trunk is the plan's own vertical circulation route: the widest fully
        # empty column band that has racking on both sides of it. (Taking merely
        # the widest empty band would pick the clear floor along a wall, and every
        # aisle would then be dragged straight through the blocks in between.)
        occupied_cols = {c for _, c in occupied}
        trunk_band = _widest_interior_band(
            [c for c in range(cmin, cmax + 1) if not any((r, c) in occupied for r in range(rmin, rmax + 1))],
            occupied_cols)
        trunk_col = trunk_band[len(trunk_band) // 2]
        trunk_x = (x_left(trunk_band[0]) + x_right(trunk_band[-1])) / 2

        h_cor = ns_aisles(locations, blockers)
        for h in h_cor:
            h["y"] = (y_top(h["rows"][0]) + y_bot(h["rows"][-1])) / 2
            h["x0"], h["x1"] = x_mid(h["cols"][0]), x_mid(h["cols"][-1])
            h["xs"], h["joined"] = set(), False
        v_cor = ew_aisles(locations)
        for v in v_cor:
            v["x"] = (x_left(v["cols"][0]) + x_right(v["cols"][-1])) / 2
            v["ys"] = {y_top(v["rows"][0]), y_bot(v["rows"][-1])}

        trunk_ys = set()
        for h in sorted(h_cor, key=lambda h: -h["y"]):
            c0, c1 = h["cols"][0], h["cols"][-1]
            # Reach the trunk, from whichever side it lies on - but only if the
            # corridor's own rows are clear all the way there.
            if trunk_col < c0 and clear_between_cols(h["rows"], c0, trunk_col):
                h["x0"], h["joined"] = trunk_x, True
                trunk_ys.add(h["y"])
            elif trunk_col > c1 and clear_between_cols(h["rows"], c1, trunk_col):
                h["x1"], h["joined"] = trunk_x, True
                trunk_ys.add(h["y"])
            elif c0 < trunk_col < c1:  # labelled on both sides: it crosses the trunk
                h["xs"].add(trunk_x)
                h["joined"] = True
                trunk_ys.add(h["y"])
            # Reach any vertical service corridor further east, same condition.
            for v in v_cor:
                if v["cols"][0] > c1 and clear_between_cols(h["rows"], c1, v["cols"][0]):
                    h["x1"] = max(h["x1"], v["x"])
                    v["ys"].add(h["y"])

        # An aisle that can't reach the trunk along its own line (10H's short
        # east end of aisle A, behind racking) gets a spur to the nearest aisle
        # that does, across clear floor, which carries on to meet it.
        def clear(rows, cols):
            return not any((r, c) in blockers for r in rows for c in cols)

        progress = True
        while progress:
            progress = False
            for h in h_cor:
                if h["joined"]:
                    continue
                best = None
                for g in h_cor:
                    if not g["joined"]:
                        continue
                    if g["rows"][-1] < h["rows"][0]:
                        between = range(g["rows"][-1] + 1, h["rows"][0])
                    elif g["rows"][0] > h["rows"][-1]:
                        between = range(h["rows"][-1] + 1, g["rows"][0])
                    else:
                        continue
                    for sc in range(h["cols"][0], h["cols"][-1] + 1):
                        if not clear(between, range(sc - 1, sc + 2)):
                            continue
                        g0, g1 = g["cols"][0], g["cols"][-1]
                        if sc > g1 and not clear(g["rows"], range(g1 + 1, sc + 1)):
                            continue
                        if sc < g0 and not clear(g["rows"], range(sc, g0)):
                            continue
                        cost = len(between) + min(abs(sc - g0), abs(sc - g1))
                        if best is None or cost < best[0]:
                            best = (cost, g, sc)
                if best:
                    _, g, sc = best
                    x = x_mid(sc)
                    h["xs"].add(x)
                    g["xs"].add(x)
                    g["x0"], g["x1"] = min(g["x0"], x), max(g["x1"], x)
                    paths.append({"id": f"spur-{h['id']}", "points": [(x, h["y"]), (x, g["y"])]})
                    h["joined"] = progress = True

        for h in h_cor:
            xs_ = sorted({h["x0"], h["x1"], *h["xs"]})
            paths.append({"id": h["id"], "points": [(x, h["y"]) for x in xs_]})

        dock_x1 = max(x_mid(h["cols"][-1]) for h in h_cor)
        # The trunk carries on through both walls when there is a building to
        # link to: the plant's buildings are linked along their central aisles.
        trunk_points = (
            ([(trunk_x, height)] if north_link_x is not None else [])
            + [(trunk_x, y) for y in sorted(trunk_ys, reverse=True)]
            + [(trunk_x, dock_y)]
            + ([(trunk_x, 0.0)] if south_link else [])
        )
        paths.append({"id": f"trunk-{building_id}", "points": trunk_points})
        paths.append({"id": f"dock-{building_id}", "points": [(trunk_x, dock_y), (dock_x1, dock_y)]})
        for v in v_cor:
            pts = sorted(v["ys"], reverse=True)
            if len(pts) > 1:
                paths.append({"id": v["id"], "points": [(v["x"], y) for y in pts]})
        if north_link_x is not None:
            north_door = (trunk_x, height)
        if south_link:
            south_door = (trunk_x, 0.0)
        pad_x = trunk_x
    else:
        # --- aisles run north-south (16G) ---------------------------------------
        def clear_between_rows(cols, ra, rb):
            lo, hi = sorted((ra, rb))
            return not any((r, c) in blockers for c in cols for r in range(lo + 1, hi))

        trunk_y = (y_top(band_rows[0]) + y_bot(band_rows[-1])) / 2
        v_cor = ew_aisles(locations)
        trunk_xs, dock_xs = set(), set()
        for v in v_cor:
            span = range(v["cols"][0], v["cols"][-1] + 1)
            v["span"] = span
            v["x"] = (x_left(span[0]) + x_right(span[-1])) / 2
            v["ys"] = {y_top(v["rows"][0]), y_bot(v["rows"][-1])}
            top, bottom = v["rows"][0], v["rows"][-1]
            if bottom < band_rows[0] and clear_between_rows(span, bottom, band_rows[0]):
                v["ys"].add(trunk_y)
                trunk_xs.add(v["x"])
            elif top > band_rows[-1] and clear_between_rows(span, band_rows[-1], top):
                v["ys"].add(trunk_y)
                trunk_xs.add(v["x"])
            elif top < band_rows[0] and bottom > band_rows[-1]:
                v["ys"].add(trunk_y)
                trunk_xs.add(v["x"])
            if bottom > band_rows[-1] and clear_between_rows(span, bottom, last_rack + 1):
                v["ys"].add(dock_y)
                dock_xs.add(v["x"])

        if north_link_x is not None:
            # Entered from the north down whichever aisle opens onto the north
            # wall nearest the building above's door.
            open_north = [v for v in v_cor if clear_between_rows(v["span"], rmin - 1, v["rows"][0])]
            entry = min(open_north, key=lambda v: abs(v["x"] - north_link_x))
            entry["ys"].add(height)
            north_door = (entry["x"], height)

        pad_x = min(dock_xs)
        if south_link:
            south_door = (pad_x, 0.0)
            paths.append({"id": f"exit-{building_id}", "points": [(pad_x, dock_y), (pad_x, 0.0)]})
        for v in v_cor:
            pts = sorted(v["ys"], reverse=True)
            paths.append({"id": v["id"], "points": [(v["x"], y) for y in pts]})
        paths.append({"id": f"trunk-{building_id}", "points": [(x, trunk_y) for x in sorted(trunk_xs)]})
        dock_end = max(max(dock_xs), pad_x + 18 + PAD_W / 2)
        paths.append({"id": f"dock-{building_id}", "points": [(x, dock_y) for x in sorted(dock_xs | {dock_end})]})

    for p in paths:
        p["buildingIds"] = [building_id]
    paths = [p for p in paths if len(p["points"]) > 1]

    # Pads sit south of the dock corridor, beside where the trunk (or, north-
    # south, the westernmost aisle) meets it.
    pad_y = dock_y - PATH_W / 2 - PAD_GAP - PAD_D / 2
    return {
        "id": building_id,
        "width": width,
        "height": height,
        "slots": slots,
        "paths": paths,
        "zones": zones_out,
        "north_door": north_door,
        "south_door": south_door,
        "lift": (pad_x + 6, pad_y),
        "delivery": (pad_x + 18, pad_y),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="data/CML_warehouse.xlsx")
    ap.add_argument("--sheet", default="PLAN_Sortie")
    ap.add_argument("--building", nargs="+", default=DEFAULT_BUILDINGS,
                    help="sheet blocks to import, north to south")
    ap.add_argument("--id", default="cml")
    ap.add_argument("--name", default="CML")
    ap.add_argument("--out", default="schema/CML.plan.json")
    ap.add_argument("--content-out", default="schema/CML.content.json")
    args = ap.parse_args()

    root = Path(__file__).resolve().parent.parent
    wb = openpyxl.load_workbook(root / args.workbook, data_only=False)
    ws = wb[args.sheet]
    sheet_zones = read_zones(ws, root / args.workbook, args.sheet)

    buildings = []
    for i, name in enumerate(args.building):
        above = buildings[-1] if buildings else None
        buildings.append(import_building(
            ws, name, sheet_zones,
            north_link_x=above["south_door"][0] if above else None,
            south_link=i < len(args.building) - 1,
        ))

    # Stack north to south: the first building's south wall stays at y = 0
    # (so a single-building import lands where it always did), each next one
    # ROAD_GAP further south.
    offset = 0.0
    for i, b in enumerate(buildings):
        if i:
            offset -= ROAD_GAP + b["height"]
        b["y0"] = offset

    walls, doors, paths, lifts, deliveries, zones, slots = [], [], [], [], [], [], []
    for i, b in enumerate(buildings):
        dy = b["y0"]
        walls.append({
            "id": b["id"],
            "closed": True,
            "points": [(0.0, dy), (b["width"], dy), (b["width"], dy + b["height"]), (0.0, dy + b["height"])],
        })
        slots.extend({**s, "y": s["y"] + dy} for s in b["slots"])
        paths.extend({**p, "points": [(x, y + dy) for x, y in p["points"]]} for p in b["paths"])
        zones.extend({**z, "y": z["y"] + dy, "buildingId": b["id"]} for z in b["zones"])
        for side, door in (("N", b["north_door"]), ("S", b["south_door"])):
            if door:
                doors.append({"id": f"Door-{b['id']}-{side}", "point": (door[0], door[1] + dy), "buildingId": b["id"]})
        lifts.append({"id": f"CL{i + 1:02d}", "point": (b["lift"][0], b["lift"][1] + dy), "buildingId": b["id"]})
        deliveries.append({"id": f"DS{i + 1:02d}", "point": (b["delivery"][0], b["delivery"][1] + dy), "buildingId": b["id"]})

    # The links: from each building's south door straight across to the north
    # door of the one below, with a jog halfway when their central aisles
    # don't line up. Their ends are the doors' own points, which the trunks
    # inside end at - the exact-coincident-point rule that joins the network.
    for north, south in zip(buildings, buildings[1:]):
        a = (north["south_door"][0], north["south_door"][1] + north["y0"])
        b = (south["north_door"][0], south["north_door"][1] + south["y0"])
        mid = (a[1] + b[1]) / 2
        points = [a, b] if abs(a[0] - b[0]) < 1e-9 else [a, (a[0], mid), (b[0], mid), b]
        paths.append({
            "id": f"link-{north['id']}-{south['id']}",
            "points": points,
            "buildingIds": [north["id"], south["id"]],
            "endpointBuildingIds": [north["id"], south["id"]],
        })

    def r3(v):
        return round(v, 3)

    def fmt(pts):
        return [{"x": r3(x), "y": r3(y)} for x, y in pts]

    def placed(items):
        return [{"id": f["id"], "x": r3(f["point"][0]), "y": r3(f["point"][1]), **({"rotationDeg": 0} if "Door" in f["id"] else {}),
                 "buildingId": f["buildingId"]} for f in items]

    config = {
        "id": args.id,
        "name": args.name,
        "units": "m",
        "walls": [{**w, "points": fmt(w["points"])} for w in walls],
        "doors": placed(doors),
        # Key order as the app's own files have it: id, points, width, buildingIds.
        "paths": [
            {"id": p["id"], "points": fmt(p["points"]), "width": PATH_W, "buildingIds": p["buildingIds"],
             **({"endpointBuildingIds": p["endpointBuildingIds"]} if "endpointBuildingIds" in p else {})}
            for p in paths
        ],
        "liftStations": placed(lifts),
        "deliverySpaces": placed(deliveries),
        "inaccessibleZones": [
            {"id": z["id"], "label": z["label"], "x": r3(z["x"]), "y": r3(z["y"]), "width": r3(z["width"]),
             "depth": r3(z["depth"]), "buildingId": z["buildingId"]}
            for z in zones
        ],
        "slotDefaults": {"width": LANE_W, "height": POS_D},
        "slots": [{**s, "x": r3(s["x"]), "y": r3(s["y"])} for s in sorted(slots, key=lambda s: s["id"])],
    }
    if not config["inaccessibleZones"]:
        del config["inaccessibleZones"]

    ids = [s["id"] for s in config["slots"]]
    if len(ids) != len(set(ids)):
        raise SystemExit("duplicate slot ids across buildings")

    out = root / args.out
    out.write_text(json.dumps(config, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {args.out}  ({len(buildings)} buildings, {len(slots)} slots, "
          f"{len(config['paths'])} paths, {len(zones)} zones)")

    # Explicit rather than derived from the plan's filename: deriving it once
    # silently pointed the empty content file at the plan itself.
    content_path = root / args.content_out
    content_path.write_text(
        json.dumps({"warehouseId": args.id, "slots": []}, indent=2) + "\n", encoding="utf-8"
    )
    print(f"wrote {args.content_out}  (empty stock; capacity is depth x {LEVELS} tiers)")


if __name__ == "__main__":
    main()
